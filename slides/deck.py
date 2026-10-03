# -*- coding: utf-8 -*-
"""
The defense deck itself. Helpers and palette live in build_slides.py.

    ../venv/Scripts/python.exe deck.py

Writes FYP_Defense_M2.pptx and FYP_Defense_ULFG.pptx -- identical content,
differing only in the second logo and the institution block.
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from build_slides import (
    blank, tb, run, para, rect, logo, title_of, takeaway, stat, bullets, table,
    figure, footer, divider,
    NAVY, DEEP, INK, GREY, MUTE, TINT, LINE, WHITE, GREEN, RED, HEAD, BODY,
    W, H, M, TY, CY, DEFENSE_DATE, VARIANTS,
)

WS_LOGO = "whitestork_software_solutions_logo.jpg"


# ============================================================ TITLE =========
def s_title(prs, v):
    s = blank(prs)
    x = M
    for stem, h in (("logo_lu.png", 0.95), (v["second_logo"], 1.06), (WS_LOGO, 0.88)):
        y = 0.52 + (1.06 - h) / 2
        x += logo(s, stem, x, y, h) + 0.5

    tf = tb(s, W - M - 2.6, 0.64, 2.6, 0.3, align=PP_ALIGN.RIGHT)
    run(para(tf, True), DEFENSE_DATE, size=12, color=GREY, italic=True)

    tf = tb(s, M, 2.30, W - 2 * M, 1.7, align=PP_ALIGN.CENTER)
    p = para(tf, True); p.alignment = PP_ALIGN.CENTER; p.line_spacing = 1.08
    run(p, "PARAMETER-EFFICIENT ADAPTATION OF A\nSMALL LANGUAGE MODEL FOR ARABIC",
        size=34, color=INK, bold=True, font=HEAD)

    ln = rect(s, 4.15, 4.12, 5.03, 0.022, fill=LINE, shape=MSO_SHAPE.RECTANGLE)
    ln.shadow.inherit = False

    tf = tb(s, M, 4.34, W - 2 * M, 0.34, align=PP_ALIGN.CENTER)
    p = para(tf, True); p.alignment = PP_ALIGN.CENTER
    run(p, "Word Sense Disambiguation  ·  Lebanese Legal Continued Pre-Training",
        size=15, color=NAVY)

    tf = tb(s, M, 4.80, W - 2 * M, 0.32, align=PP_ALIGN.CENTER)
    p = para(tf, True); p.alignment = PP_ALIGN.CENTER
    run(p, v["banner"], size=12, color=GREY, bold=True)

    tf = tb(s, M, 5.80, 6.2, 1.2)
    p = para(tf, True); p.space_after = Pt(4)
    run(p, "Presented by  ", size=12.5, color=GREY)
    run(p, "Rokaya Al Harakeh", size=12.5, color=INK, bold=True)
    p = para(tf); p.space_after = Pt(3)
    run(p, "Academic supervisor:  Dr. [ Name ]", size=11.5, color=GREY)
    p = para(tf)
    run(p, "Company supervisor:  Dr. Daoud Baalbaki  —  WhiteStork", size=11.5, color=GREY)

    tf = tb(s, W - M - 5.3, 5.80, 5.3, 1.2, align=PP_ALIGN.RIGHT)
    for i, line in enumerate([t for t in v["institutions"] if t]):
        p = para(tf, i == 0); p.alignment = PP_ALIGN.RIGHT; p.space_after = Pt(2)
        run(p, line, size=11.5, color=GREY)
    return s


# ============================================================ CONTEXT ======
def s_objectives(prs):
    s = blank(prs)
    title_of(s, "Adapting one small model under two objectives", kicker="Context")
    tf = tb(s, M, CY, 6.2, 1.9)
    p = para(tf, True); p.space_after = Pt(8)
    run(p, "Word Sense Disambiguation", size=16, color=INK, bold=True)
    p = para(tf); p.space_after = Pt(10)
    run(p, "Given a sentence, a target word and candidate definitions, choose the "
           "intended sense. Arabic makes this harder: rich morphology, attached "
           "clitics, and diacritics omitted in ordinary writing.", size=14, color=GREY)
    p = para(tf)
    run(p, "Large models are expensive to adapt. Can a 2B model compete?",
        size=14.5, color=DEEP, bold=True)

    rows = [["", "Objective 1 — SFT", "Objective 2 — CPT"],
            ["Data", "9,952 labelled items", "27.2M tokens of legal text"],
            ["Signal", "gold sense IDs", "none — raw text"],
            ["Compared to", "a published paper", "the untrained base model"]]
    table(s, 7.05, CY - 0.04, 5.66, rows, [1.32, 2.0, 2.34], size=12, row_h=0.46)

    tfh = tb(s, M, 3.50, W - 2 * M, 0.32)
    run(para(tfh, True), "Four objectives", size=16, color=INK, bold=True)
    bullets(s, M, 3.88, W - 2 * M, [
        ("Reproduce at 2B. ", "Match the published Gemma 2-9B result with a model "
         "4.5× smaller, on a free GPU."),
        ("Establish the floor. ", "The literature on this benchmark reports no "
         "baselines. Establish them."),
        ("Audit the metrics. ", "Determine whether the reported scores carry the "
         "information they are assumed to."),
        ("Adapt to a domain. ", "Continued pre-training on a Lebanese legal corpus, "
         "contrasted against SFT."),
    ], size=14, gap=5, h=2.2)
    takeaway(s, "The question is never which objective is better — it is "
                "**what each one changes in the model**.")
    return s


def s_finetuning(prs):
    s = blank(prs)
    title_of(s, "Why fine-tuning a 2B model needs a trick", kicker="Fine-tuning")
    tf = tb(s, M, CY, 5.9, 1.3)
    p = para(tf, True); p.space_after = Pt(8)
    run(p, "Full fine-tuning updates every weight. The weights are not the "
           "bottleneck — the optimiser state is.", size=15, color=GREY)
    p = para(tf)
    run(p, "Adam keeps two moments per trainable parameter, in fp32.",
        size=14, color=GREY, italic=True)

    stat(s, M, 2.86, 2.8, "41.7 GB", "Full fine-tuning", "training state needed",
         color=RED, h=1.55)
    stat(s, M + 3.0, 2.86, 2.8, "16 GB", "Available", "free-tier Colab T4",
         color=GREY, h=1.55)

    tfa = tb(s, M, 4.72, 5.8, 1.5)
    p = para(tfa, True)
    run(p, "Only 5.2 GB of that is weights. ", size=14.5, color=GREY)
    run(p, "36.5 GB is gradients and optimiser state — which scales with the number "
           "of trainable parameters, not with model size.", size=14.5, color=INK, bold=True)

    figure(s, "fig_memory.png", 7.05, 1.78, w=5.7)
    takeaway(s, "So freezing the base is not one option among several. "
                "**It is the only one that fits.**")
    return s


def s_lora(prs):
    s = blank(prs)
    title_of(s, "LoRA and QLoRA, in one slide", kicker="Method")
    tf = tb(s, M, CY, 6.0, 2.3)
    p = para(tf, True); p.space_after = Pt(10)
    run(p, "LoRA.  ", size=15, color=INK, bold=True)
    run(p, "Freeze W₀ and learn a low-rank update B·A. Fine-tuning updates are known "
           "to have low intrinsic rank, so little is given up.", size=15, color=GREY)
    p = para(tf); p.space_after = Pt(10)
    run(p, "QLoRA.  ", size=15, color=INK, bold=True)
    run(p, "Hold the frozen base in 4-bit NF4 with double quantization and compute in "
           "bfloat16. The adapter itself stays full precision.", size=15, color=GREY)
    p = para(tf)
    run(p, "It merges back into the weights after training, so unlike adapters or "
           "prefix tuning it adds no inference cost.", size=14, color=DEEP, italic=True)

    stat(s, M, 4.28, 2.8, "1.6%", "of parameters trained", "41.5M of 2.61B", h=1.5)
    stat(s, M + 3.0, 4.28, 2.8, "16×", "less training memory", "versus full fine-tuning",
         color=GREEN, h=1.5)
    figure(s, "fig_lora.png", 7.1, 1.7, w=5.65)
    takeaway(s, "Both halves of this project are QLoRA runs on the same base model. "
                "**Only the training objective changes.**")
    return s


def s_two(prs):
    s = blank(prs)
    title_of(s, "Two objectives that change different things",
             kicker="The central distinction")
    rows = [["", "Supervised fine-tuning", "Continued pre-training"],
            ["Learns from", "instruction → response pairs", "raw text, next-token prediction"],
            ["Supervision", "gold labels", "none"],
            ["Sequence unit", "one example per sequence", "packed 2,048-token blocks"],
            ["LoRA rank · epochs", "32  ·  3", "16  ·  1"],
            ["Teaches the model", "how to answer", "what the domain looks like"]]
    table(s, M, CY, W - 2 * M, rows, [2.6, 4.75, 4.74], size=13.5, row_h=0.5,
          bold_rows=(5,))
    tf = tb(s, M, 5.12, W - 2 * M, 1.1)
    p = para(tf, True); p.space_after = Pt(7)
    run(p, "The last row is the thesis. ", size=15.5, color=INK, bold=True)
    run(p, "SFT teaches a task format; CPT teaches a domain distribution. The rest of "
           "this talk measures both, on the same model, to show the difference is real "
           "and not merely plausible.", size=15, color=GREY)
    takeaway(s, "They are not substitutes. The choice follows from "
                "**what is missing — format, or knowledge**.")
    return s


# ============================================================ PART 1 =======
def s_sft_setup(prs):
    s = blank(prs)
    title_of(s, "The task, the data, the run", kicker="Part 1 · SFT")
    table(s, M, CY, 5.7, [["Dataset A — El-Razzaz", ""],
                          ["Train / dev / test", "9,952 / 2,487 / 3,110"],
                          ["Candidates per item", "exactly 2 glosses"],
                          ["Output", "the correct sense ID"]],
          [3.2, 2.5], size=13, row_h=0.46)
    table(s, 7.05, CY, 5.68, [["Configuration", ""],
                              ["Base model", "Gemma 2-2B, 4-bit NF4"],
                              ["LoRA r / α / dropout", "32 / 32 / 0.05"],
                              ["Epochs · LR", "3  ·  2e-4, adamw_8bit"],
                              ["Hardware", "free Colab T4, 16 GB"]],
          [2.86, 2.82], size=13, row_h=0.46)

    tf = tb(s, M, 4.54, W - 2 * M, 1.7)
    p = para(tf, True); p.space_after = Pt(9)
    run(p, "Everything the model needs is already in the prompt. ", size=15.5,
        color=INK, bold=True)
    run(p, "The sentence, the target word and both candidate definitions are all "
           "supplied. Nothing has to be recalled from the weights — the model only "
           "has to choose between two visible strings.", size=15, color=GREY)
    p = para(tf)
    run(p, "That single observation is what the whole of Part 1 ends up explaining.",
        size=14.5, color=DEEP, italic=True)
    takeaway(s, "A 2B model on a **free** GPU, against published 7–9B models trained "
                "on an NVIDIA L4.")
    return s


def s_sft_headline(prs):
    s = blank(prs)
    title_of(s, "90.42% — matching models four times larger", kicker="Part 1 · Result")
    stat(s, M, CY, 2.72, "90.42%", "Accuracy", "2,812 of 3,110 correct", h=1.5)
    stat(s, M + 2.92, CY, 2.72, "0.8333", "Macro-F1", "but see the next slide", h=1.5)
    stat(s, M + 5.84, CY, 2.72, "0", "Malformed outputs", "across the whole test set",
         color=GREEN, h=1.5)

    table(s, M, 3.36, 8.56, [["Model", "Params", "Accuracy", "Macro-F1"],
                             ["Gemma 2-9B (published)", "9B", "89.39", "81.72"],
                             ["LLaMA 3.1-8B (published)", "8B", "90.42", "83.20"],
                             ["Qwen 2.5-7B (published)", "7B", "90.77", "83.98"],
                             ["Gemma 2-2B (this work)", "2B", "90.42", "83.33"]],
          [3.46, 1.5, 1.8, 1.8], size=13, row_h=0.42, bold_rows=(4,))

    card = rect(s, 9.86, 3.36, 2.87, 2.52, fill=WHITE, line=LINE)
    tfc = card.text_frame
    tfc.margin_left = tfc.margin_right = Inches(0.18); tfc.margin_top = Inches(0.16)
    p = tfc.paragraphs[0]; p.space_after = Pt(7)
    run(p, "Is the difference real?", size=14.5, color=INK, bold=True)
    p = tfc.add_paragraph(); p.space_after = Pt(7)
    run(p, "SE = 0.53 pp at n = 3,110\n95% CI ≈ [89.4, 91.4]", size=13, color=GREY)
    p = tfc.add_paragraph()
    run(p, "Every published result falls inside it. Differences under one point are "
           "not distinguishable.", size=12.5, color=RED)

    takeaway(s, "The defensible claim is not that 2B is better. It is that it is "
                "**not worse**, at a quarter of the parameters and none of the cost.")
    return s


def s_sft_baselines(prs):
    s = blank(prs)
    title_of(s, "The benchmark's real floor is not 50%", kicker="Part 1 · Contribution")
    figure(s, "fig_baselines.png", M, CY + 0.12, w=6.85)
    tf = tb(s, 7.8, CY, 4.93, 3.6)
    p = para(tf, True); p.space_after = Pt(11)
    run(p, "Every item offers two candidates, so random choice scores 50%. But the "
           "gold answer sits in second position ", size=14.5, color=GREY)
    run(p, "64.95%", size=14.5, color=RED, bold=True)
    run(p, " of the time.", size=14.5, color=GREY)
    p = para(tf); p.space_after = Pt(11)
    run(p, "A one-line program that always picks the second candidate — no model, "
           "no Arabic — scores 64.95%.", size=14.5, color=INK, bold=True)
    p = para(tf); p.space_after = Pt(11)
    run(p, "It is stable across splits (65.74 / 64.94 / 64.95), so it is a property "
           "of how the dataset was constructed.", size=14, color=GREY)
    p = para(tf)
    run(p, "Classical Lesk adds essentially nothing over it.", size=14, color=GREY)
    takeaway(s, "No baseline appears anywhere in the published literature on this "
                "dataset. The real gain is **90.42 over 64.95**, not over 50.")
    return s


def s_sft_metric(prs):
    s = blank(prs)
    title_of(s, "The reported secondary metric is degenerate",
             kicker="Part 1 · Contribution")
    tf = tb(s, M, CY, 6.4, 2.9)
    p = para(tf, True); p.space_after = Pt(10)
    run(p, "Every gold class has a support of exactly one — 3,110 instances, 3,110 "
           "distinct gold labels.", size=15, color=GREY)
    p = para(tf); p.space_after = Pt(12)
    run(p, "Macro-averaging takes the union of true and predicted labels, so a class "
           "appearing only in the predictions scores zero.", size=15, color=GREY)
    p = para(tf)
    run(p, "macro-recall   =   correct  /  | y_true ∪ y_pred |",
        size=16.5, color=DEEP, bold=True, font=HEAD)

    card = rect(s, 7.2, CY - 0.04, 5.53, 2.96, fill=WHITE, line=LINE)
    tfc = card.text_frame
    tfc.margin_left = tfc.margin_right = Inches(0.22); tfc.margin_top = Inches(0.18)
    p = tfc.paragraphs[0]; p.space_after = Pt(8)
    run(p, "Verify it arithmetically", size=15, color=INK, bold=True)
    for line, col, bold in [("0.9042 × 3,110  =  2,812 correct", GREY, False),
                            ("2,812 / 0.8379  =  3,356.0  exactly", GREY, False),
                            ("3,356 − 3,110  =  246 phantom classes", RED, True)]:
        p = tfc.add_paragraph(); p.space_after = Pt(7)
        run(p, line, size=14, color=col, bold=bold, font=HEAD)
    p = tfc.add_paragraph()
    run(p, "Those 246 are created by the model's own wrong answers — the more varied "
           "its errors, the larger the denominator.", size=12.5, color=GREY)

    takeaway(s, "Macro-F1 here is **accuracy rescaled by a denominator the model "
                "inflates**. It carries no independent information — and the "
                "literature reports it as though it does.")
    return s


def s_sft_errors(prs):
    s = blank(prs)
    title_of(s, "What the 298 errors are made of", kicker="Part 1 · Error analysis")
    figure(s, "fig_errors.png", M, CY + 0.1, w=6.8)
    tf = tb(s, 7.75, CY, 4.98, 3.7)
    for i, (lead, rest) in enumerate([
        ("Zero parse failures. ", "All 3,110 outputs were valid sense IDs. The errors "
         "are genuine confusions, not formatting."),
        ("No positional bias. ", "Despite the 64.95% second-position skew, errors split "
         "evenly — the model learned the task, not the shortcut."),
        ("Close glosses dominate. ", "Errors concentrate where the two definitions share "
         "vocabulary. Some pairs are near-duplicates that cap the achievable ceiling."),
    ]):
        p = para(tf, i == 0); p.space_after = Pt(13)
        run(p, lead, size=14.5, color=INK, bold=True)
        run(p, rest, size=14.5, color=GREY)
    takeaway(s, "A 2B model matches 9B because this task is **discrimination between "
                "two visible strings, not knowledge retrieval**. Capacity buys knowledge; "
                "knowledge is not what is short here.")
    return s
