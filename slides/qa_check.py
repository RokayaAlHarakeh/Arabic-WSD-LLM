# -*- coding: utf-8 -*-
"""
Geometric QA for the decks, since no renderer is installed on this machine.

Catches the defects a visual pass would catch first:
  * anything off-slide or inside the 0.5" margin
  * text that will not fit its box at its own font size
  * content-bearing shapes that overlap each other

Estimates are deliberately conservative: a flag is a thing to look at, not proof.

    ../venv/Scripts/python.exe qa_check.py FYP_Defense_M2.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0
W, H = 13.333, 7.5
MARGIN = 0.5
# average glyph width as a fraction of point size, for a proportional sans
CHAR_W = 0.50
LINE_H = 1.22


def inches(v):
    return (v or 0) / EMU


def text_lines(tf, box_w):
    """Height the text needs, in inches, summed per paragraph at its own size.

    Summing per paragraph matters: a stat card is one 40pt line over two small
    ones, and multiplying every line by the largest size overstates it by 2x.
    """
    need, biggest, lines = 0.0, 0.0, 0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        size = max([(r.font.size.pt if r.font.size else 18.0) for r in p.runs] or [18.0])
        biggest = max(biggest, size)
        after = p.space_after.pt if p.space_after is not None else 0.0
        if not txt.strip():
            need += size * LINE_H / 72.0
            lines += 1
            continue
        n = 0
        for hard in txt.split("\n"):
            cpl = max(1, int(box_w * 72.0 / (size * CHAR_W)))
            n += max(1, -(-len(hard) // cpl))
        lines += n
        need += (n * size * LINE_H + after) / 72.0
    return need, biggest, lines


def overlaps(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    ix = min(ax + aw, bx + bw) - max(ax, bx)
    iy = min(ay + ah, by + bh) - max(ay, by)
    if ix <= 0.02 or iy <= 0.02:
        return 0.0
    return ix * iy


def check(path):
    prs = Presentation(path)
    print("=" * 70)
    print(path)
    print("=" * 70)
    total = 0
    for idx, slide in enumerate(prs.slides, 1):
        issues = []
        boxes = []
        for sh in slide.shapes:
            x, y = inches(sh.left), inches(sh.top)
            w, h = inches(sh.width), inches(sh.height)
            name = sh.shape_type
            # full-bleed backgrounds are intentional
            full_bleed = w > W - 0.1 and h > H - 0.1

            if not full_bleed:
                if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
                    issues.append("OFF-SLIDE   %-18s x=%.2f y=%.2f w=%.2f h=%.2f"
                                  % (str(name)[:18], x, y, w, h))
                elif x < MARGIN - 0.02 or y < MARGIN - 0.02 \
                        or x + w > W - MARGIN + 0.02 or y + h > H - MARGIN + 0.02:
                    # footer and page number sit low on purpose
                    if y < H - 0.62:
                        issues.append("TIGHT MARGIN %-17s x=%.2f y=%.2f r=%.2f b=%.2f"
                                      % (str(name)[:17], x, y, x + w, y + h))

            if sh.has_text_frame and sh.text_frame.text.strip():
                inner = w - (inches(sh.text_frame.margin_left)
                             + inches(sh.text_frame.margin_right))
                need, size, n = text_lines(sh.text_frame, max(inner, 0.4))
                need += inches(sh.text_frame.margin_top) + inches(sh.text_frame.margin_bottom)
                if need > h + 0.08:
                    issues.append("TEXT OVERFLOW  needs %.2f\" has %.2f\"  (%d lines, max %.0fpt)"
                                  "  %r" % (need, h, n, size,
                                            sh.text_frame.text[:48].replace("\n", " ")))
                if size < 10:
                    issues.append("TINY TEXT  %.0fpt  %r" % (size, sh.text_frame.text[:40]))
                boxes.append(((x, y, w, h), sh.text_frame.text[:28].replace("\n", " ")))

        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                area = overlaps(a[0], b[0])
                small = min(a[0][2] * a[0][3], b[0][2] * b[0][3])
                if area > 0.3 * small:
                    issues.append("OVERLAP  %r  <>  %r" % (a[1], b[1]))

        if issues:
            total += len(issues)
            print("\nslide %d" % idx)
            for m in issues:
                print("   " + m)
    print("\n%s  %d issue(s)" % ("CLEAN" if not total else "REVIEW", total))
    return total


if __name__ == "__main__":
    files = sys.argv[1:] or ["FYP_Defense_M2.pptx"]
    for f in files:
        check(f)
