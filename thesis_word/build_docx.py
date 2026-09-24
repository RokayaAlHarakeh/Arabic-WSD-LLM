# -*- coding: utf-8 -*-
"""
Build the FYP report as a Microsoft Word .docx.

    python build_docx.py

Produces  FYP_Report.docx  in this folder.

Design notes
------------
* Real Word Heading 1/2/3 styles are used, so the Navigation pane and the
  automatic Table of Contents both work.
* Arabic runs are explicitly marked right-to-left with a complex-script font,
  which is what Word needs to render and shape Arabic correctly.
* Tables use a booktabs-like rule scheme (top rule, header rule, bottom rule,
  no vertical lines) rather than Word's default grid.
* Equations are inserted as centred Unicode text. Word's equation editor
  cannot be driven from python-docx; the handful that matter are flagged
  with EQ-TODO in the document so they can be replaced by hand.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT  = os.path.join(os.path.dirname(os.path.abspath(__file__)), "FYP_Report.docx")
OUT2 = os.path.join(os.path.dirname(os.path.abspath(__file__)), "FYP_Report_Engineering.docx")

# Fonts -----------------------------------------------------------------
LATIN_FONT  = "Times New Roman"
ARABIC_FONT = "Traditional Arabic"   # on Windows. Alternatives: "Amiri", "Simplified Arabic"
MONO_FONT   = "Consolas"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # dark blue, for headings
TODOCOL = RGBColor(0xB0, 0x00, 0x00)  # red, for TODO markers


# ======================================================================
#  Low-level helpers
# ======================================================================
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def mark_rtl(run, font=ARABIC_FONT, size=None):
    """Make a run render as right-to-left Arabic."""
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = _el('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:cs'), font)
    rFonts.set(qn('w:ascii'), font)
    rFonts.set(qn('w:hAnsi'), font)
    rPr.append(_el('w:rtl'))
    if size:
        rPr.append(_el('w:szCs', **{'w:val': str(int(size * 2))}))


def add_field(paragraph, instruction, placeholder=""):
    """Insert a Word field (used for TOC and PAGE numbers)."""
    r1 = paragraph.add_run()
    r1._element.append(_el('w:fldChar', **{'w:fldCharType': 'begin'}))
    r2 = paragraph.add_run()
    it = _el('w:instrText', **{'xml:space': 'preserve'})
    # Word wants the field code padded with spaces inside the delimiters. Without
    # them some builds fail to parse the code and the field never recalculates —
    # which is what left every SEQ caption frozen at its placeholder value.
    it.text = " %s " % instruction.strip()
    r2._element.append(it)
    r3 = paragraph.add_run()
    r3._element.append(_el('w:fldChar', **{'w:fldCharType': 'separate'}))
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run()
    r5._element.append(_el('w:fldChar', **{'w:fldCharType': 'end'}))


def set_page_numbering(section, fmt="decimal", start=None):
    """fmt: 'decimal' | 'lowerRoman'"""
    sectPr = section._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = _el('w:pgNumType')
        sectPr.append(pg)
    pg.set(qn('w:fmt'), fmt)
    if start is not None:
        pg.set(qn('w:start'), str(start))


def add_footer_pagenum(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE", "1")


def booktabs(table, header_rows=1):
    """Top rule, rule under header, bottom rule. No vertical lines."""
    tbl = table._tbl
    # strip any existing borders definition
    tblPr = tbl.tblPr
    for b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(b)
    borders = _el('w:tblBorders')
    for edge, sz in (('top', '12'), ('bottom', '12')):
        borders.append(_el(f'w:{edge}', **{'w:val': 'single', 'w:sz': sz,
                                           'w:space': '0', 'w:color': '000000'}))
    for edge in ('left', 'right', 'insideV'):
        borders.append(_el(f'w:{edge}', **{'w:val': 'none', 'w:sz': '0', 'w:space': '0'}))
    borders.append(_el('w:insideH', **{'w:val': 'none', 'w:sz': '0', 'w:space': '0'}))
    tblPr.append(borders)
    # rule under the header row
    for r in range(header_rows):
        for cell in table.rows[r].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcb = _el('w:tcBorders')
            tcb.append(_el('w:bottom', **{'w:val': 'single', 'w:sz': '8',
                                          'w:space': '0', 'w:color': '000000'}))
            tcPr.append(tcb)


def shade(paragraph, hexcolor="F2F2F2"):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(_el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': hexcolor}))


# ======================================================================
#  Document-level helpers
# ======================================================================
class Doc:
    def __init__(self):
        self.d = Document()
        self.seq = {}          # running Table / Figure / Listing counters
        self._setup_styles()
        self._setup_page()

    # -- styles ---------------------------------------------------------
    def _setup_styles(self):
        st = self.d.styles['Normal']
        st.font.name = LATIN_FONT
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), LATIN_FONT)
        pf = st.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.4
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for name, size, bold, before, after in (
            ('Heading 1', 18, True, 24, 12),
            ('Heading 2', 14, True, 18, 8),
            ('Heading 3', 12.5, True, 14, 6),
        ):
            s = self.d.styles[name]
            s.font.name = LATIN_FONT
            s.font.size = Pt(size)
            s.font.bold = bold
            s.font.color.rgb = ACCENT
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            s.paragraph_format.keep_with_next = True

    def _setup_page(self):
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin = Cm(3), Cm(2.5)
        s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.5)

    # -- content --------------------------------------------------------
    def h1(self, text, numbered=True):
        p = self.d.add_paragraph(text, style='Heading 1')
        p.paragraph_format.page_break_before = True
        return p

    def h2(self, text):
        return self.d.add_paragraph(text, style='Heading 2')

    def h3(self, text):
        return self.d.add_paragraph(text, style='Heading 3')

    def p(self, text=None, bold_parts=None, align=None):
        """text may contain **bold** and \\ar{arabic} markers."""
        par = self.d.add_paragraph()
        if align:
            par.alignment = align
        if text:
            self._rich(par, text)
        return par

    def _rich(self, par, text):
        """Parse **bold**, *italic*, `mono`, and \\ar{...} inline markers."""
        import re
        pattern = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\\ar\{[^}]+\})', re.S)
        for piece in pattern.split(text):
            if not piece:
                continue
            if piece.startswith('**') and piece.endswith('**'):
                par.add_run(piece[2:-2]).bold = True
            elif piece.startswith('*') and piece.endswith('*') and len(piece) > 2:
                par.add_run(piece[1:-1]).italic = True
            elif piece.startswith('`') and piece.endswith('`'):
                r = par.add_run(piece[1:-1]); r.font.name = MONO_FONT; r.font.size = Pt(10.5)
            elif piece.startswith('\\ar{'):
                r = par.add_run(piece[4:-1]); mark_rtl(r, size=13)
            else:
                par.add_run(piece)

    def bullet(self, text, level=0):
        par = self.d.add_paragraph(style='List Bullet')
        par.paragraph_format.left_indent = Cm(0.8 + 0.6 * level)
        self._rich(par, text)
        return par

    def numbered(self, text):
        par = self.d.add_paragraph(style='List Number')
        self._rich(par, text)
        return par

    def eq(self, text, note=None):
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(8)
        r = par.add_run(text)
        r.italic = True
        r.font.size = Pt(12)
        if note:
            n = self.d.add_paragraph()
            n.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = n.add_run(f"[EQ-TODO: {note}]")
            rr.font.size = Pt(8); rr.font.color.rgb = TODOCOL
        return par

    def keypoint(self, text):
        par = self.d.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.5)
        par.paragraph_format.right_indent = Cm(0.5)
        par.paragraph_format.space_before = Pt(10)
        par.paragraph_format.space_after = Pt(10)
        self._rich(par, text)
        shade(par, "EEF3F8")
        return par

    def todo(self, text):
        par = self.d.add_paragraph()
        par.add_run("TODO — ")
        self._rich(par, text)          # so \ar{} markers still render RTL
        for r in par.runs:
            r.font.color.rgb = TODOCOL
            r.font.size = Pt(10.5)
            r.italic = True
        return par

    def caption(self, text, kind="Table"):
        """A real Word caption: styled 'Caption' with a SEQ field, so that
        numbering is automatic AND the List of Tables / Figures can find it."""
        par = self.d.add_paragraph(style='Caption')
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(10)
        r = par.add_run(f"{kind} ")
        r.bold = True
        # The field's cached result is the number this caption will get anyway,
        # so the document reads correctly before anyone presses F9; updating the
        # fields in Word then recomputes the same values.
        self.seq[kind] = self.seq.get(kind, 0) + 1
        add_field(par, f'SEQ {kind} \\* ARABIC', str(self.seq[kind]))
        r2 = par.add_run(" — " + text)
        for run in par.runs:
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
        par.runs[0].bold = True
        return par

    def list_of(self, kind):
        """List of Tables / List of Figures — a TOC field filtered by caption label."""
        par = self.d.add_paragraph()
        add_field(par, f'TOC \\h \\z \\c "{kind}"',
                  f"Right-click here → Update Field to build the list of {kind.lower()}s.")
        return par

    def listing(self, lines, caption=None, arabic_lines=(), size=9.5):
        """A monospace, shaded code/data block. `arabic_lines` holds the indices
        of lines that contain Arabic and must be rendered right-to-left."""
        for i, line in enumerate(lines):
            par = self.d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            par.paragraph_format.right_indent = Cm(0.3)
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if i in arabic_lines:
                self._rich(par, line)
                for r in par.runs:
                    if any('؀' <= ch <= 'ۿ' for ch in r.text):
                        # raw Arabic (not wrapped in \ar{}) still needs RTL marking
                        rPr = r._element.find(qn('w:rPr'))
                        if rPr is None or rPr.find(qn('w:rtl')) is None:
                            mark_rtl(r, size=11)
                    else:
                        r.font.name = MONO_FONT
                    r.font.size = Pt(size)
            else:
                r = par.add_run(line if line else " ")
                r.font.name = MONO_FONT
                r.font.size = Pt(size)
            shade(par, "F4F4F4")
        if caption:
            self.caption(caption, kind="Listing")

    def figure(self, image, text, note=None, width_cm=14.0):
        """Embed figures/<image> with a numbered caption. If the file is
        missing, fall back to a visible red placeholder so nothing is lost."""
        import os as _os
        path = _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "figures", image)
        if _os.path.exists(path):
            par = self.d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(10)
            par.paragraph_format.space_after = Pt(2)
            par.add_run().add_picture(path, width=Cm(width_cm))
        else:
            box = self.d.add_paragraph()
            box.alignment = WD_ALIGN_PARAGRAPH.CENTER
            box.paragraph_format.space_before = Pt(10)
            r = box.add_run("[ FIGURE MISSING: %s%s ]" % (image, " - " + note if note else ""))
            r.font.color.rgb = TODOCOL
            r.font.size = Pt(10.5)
            r.italic = True
            shade(box, "FBEAEA")
        self.caption(text, kind="Figure")
        return self

    def figure_placeholder(self, text, note):
        """Reserve a slot for a figure that still has to be produced by hand."""
        box = self.d.add_paragraph()
        box.alignment = WD_ALIGN_PARAGRAPH.CENTER
        box.paragraph_format.space_before = Pt(10)
        r = box.add_run("[ FIGURE TO INSERT - %s ]" % note)
        r.font.color.rgb = TODOCOL
        r.font.size = Pt(10.5)
        r.italic = True
        shade(box, "FBEAEA")
        self.caption(text, kind="Figure")
        return box

    def table(self, rows, widths=None, align_right=None):
        """rows[0] is the header. Cell text supports the same inline markers."""
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, cell_text in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(3)
                par.paragraph_format.line_spacing = 1.0
                par.alignment = (WD_ALIGN_PARAGRAPH.RIGHT
                                 if align_right and j in align_right and i > 0
                                 else WD_ALIGN_PARAGRAPH.LEFT)
                txt = f"**{cell_text}**" if i == 0 and not cell_text.startswith('**') else cell_text
                self._rich(par, txt)
                for r in par.runs:
                    r.font.size = Pt(10.5)
        booktabs(t)
        if widths:
            for j, w in enumerate(widths):
                for row in t.rows:
                    row.cells[j].width = Cm(w)
        return t

    def pagebreak(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        try:
            self.d.save(path)
            return path
        except PermissionError:
            # The file is open in Word. Write beside it rather than failing.
            base, ext = os.path.splitext(path)
            alt = f"{base}_NEW{ext}"
            self.d.save(alt)
            print("!! FYP_Report.docx is open in Word, so it could not be overwritten.")
            print(f"!! Wrote {os.path.basename(alt)} instead. Close Word and rerun to replace"
                  " the original, or just use the _NEW file.")
            return alt


# ======================================================================
#  TITLE PAGE
# ======================================================================
FIG_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures")

# Details shared by both title-page variants.
AUTHOR = "Rokaya Al Harakeh"
TITLE = ("Parameter-Efficient Adaptation of a Small Language Model "
         "for Arabic Word Sense Disambiguation")


def _find_logo(stem):
    """figures/<stem>.<ext> — whichever extension was actually saved."""
    for ext in (".png", ".PNG", ".jpg", ".JPG", ".jpeg", ".JPEG"):
        p = os.path.join(FIG_DIR, stem + ext)
        if os.path.exists(p):
            return p
    return None


def _logo_height(path):
    """Height in cm that puts the *emblem* at roughly 2.4 cm.

    A logo taller than it is wide almost always carries a caption beneath the
    emblem — the Faculty of Sciences one spends its bottom 21.5% on three lines
    of text. Scaling those to the same overall height as a square emblem leaves
    them looking about a fifth too small, so give them more room and let the
    caption hang below.
    """
    try:
        from docx.image.image import Image as _Img
        im = _Img.from_file(path)
        return 3.05 if im.px_width / im.px_height < 0.95 else 2.40
    except Exception:
        return 2.40


def _borderless(table):
    tblPr = table._tbl.tblPr
    for b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(b)
    nb = _el('w:tblBorders')
    for edge in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
        nb.append(_el(f'w:{edge}', **{'w:val': 'none', 'w:sz': '0', 'w:space': '0'}))
    tblPr.append(nb)
    return table


def _logo_row(d, left_stem, right_stem):
    """University emblem left, faculty emblem right, tops aligned.

    A borderless two-cell table rather than a tab stop, because inline images
    sit on the text baseline: with unequal heights that would hang the shorter
    emblem low instead of lining the two up along their tops.
    """
    left, right = _find_logo(left_stem), _find_logo(right_stem)
    t = _borderless(d.add_table(rows=1, cols=2))
    t.autofit = False
    for cell, img, stem, al in ((t.rows[0].cells[0], left, left_stem, WD_ALIGN_PARAGRAPH.LEFT),
                                (t.rows[0].cells[1], right, right_stem, WD_ALIGN_PARAGRAPH.RIGHT)):
        cell.width = Cm(7.75)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cp = cell.paragraphs[0]
        cp.alignment = al
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        if img:
            cp.add_run().add_picture(img, height=Cm(_logo_height(img)))
        else:
            # Say which file is missing rather than silently shifting the layout,
            # so the slot is obvious and fills itself once the file is saved.
            r = cp.add_run("[ save %s.png\n  in thesis_word/figures/ ]" % stem)
            r.font.size = Pt(9)
            r.italic = True
            r.font.color.rgb = TODOCOL
    d.add_paragraph().paragraph_format.space_after = Pt(16)


def _label_block(d, rows, label_cm=3.2, value_cm=7.0):
    """A centred 'Supervisor / Reviewers' block: label column, then names."""
    flat = [(lab, n, i == 0) for lab, names in rows for i, n in enumerate(names)]
    t = _borderless(d.add_table(rows=len(flat), cols=2))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (lab, name, first) in enumerate(flat):
        for j, (txt, w, bold, al) in enumerate(
                ((lab if first else "", label_cm, True, WD_ALIGN_PARAGRAPH.LEFT),
                 (name, value_cm, False, WD_ALIGN_PARAGRAPH.LEFT))):
            cell = t.cell(i, j)
            cell.width = Cm(w)
            par = cell.paragraphs[0]
            par.alignment = al
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(3)
            par.paragraph_format.line_spacing = 1.0
            r = par.add_run(txt)
            r.font.size = Pt(12.5)
            r.bold = bold
            if txt.startswith("Dr. ["):
                r.font.color.rgb = TODOCOL
    return t


def _jury(d, rows, tab_cm=13.8):
    """Jury lines with real Word dot leaders rather than typed full stops."""
    for name, role in rows:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(tab_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(name)
        r.font.size = Pt(12)
        if "[" in name:
            r.font.color.rgb = TODOCOL
        p.add_run("\t")
        r2 = p.add_run(role)
        r2.font.size = Pt(12)


def title_page(D, variant):
    """variant: 'sciences' (Faculty of Sciences) or 'engineering' (FoE Branch III)."""
    d = D.d

    def centre(text, size=12, bold=False, color=None, space=6):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
        return p

    d.add_paragraph().paragraph_format.space_after = Pt(0)

    if variant == "engineering":
        _logo_row(d, "logo_lu", "ulfg_logo")
        centre("FINAL YEAR PROJECT REPORT", 20, True, ACCENT, 16)
        centre("Submitted in fulfillment of the requirements for the", 12, space=8)
        centre("ENGINEERING DEGREE FROM THE LEBANESE UNIVERSITY –", 13.5, True, TODOCOL, 2)
        centre("FACULTY OF ENGINEERING – BRANCH III", 13.5, True, TODOCOL, 18)
        centre("Major: Electrical and Electronics Engineering", 12.5, space=3)
        centre("Option: Computer and Telecommunications Engineering", 12.5, space=24)
        centre("Prepared by:", 12, space=6)
        centre(AUTHOR, 16, True, space=24)
        centre("Project Title:", 12, space=8)
        centre(TITLE, 17, True, TODOCOL, 26)
        centre("Supervised by:", 12, space=6)
        centre("Dr. [Academic supervisor]", 13, True, TODOCOL, 2)
        centre("Dr. Daoud Baalbaki (WhiteStork)", 13, True, space=24)
        centre("Defended on [DD / MM / 2026] in front of the jury:", 12, False, TODOCOL, 12)
        _jury(d, [("Dr. [Jury member 1]", "President"),
                  ("Dr. [Jury member 2]", "Member"),
                  ("Dr. [Jury member 3]", "Member")])
    else:
        _logo_row(d, "logo_lu", "logo_institute")
        centre("MASTER THESIS", 20, True, ACCENT, 18)
        centre("In Order to Obtain the", 12, space=8)
        centre("PROFESSIONAL MASTER", 16, True, TODOCOL, 8)
        centre("in", 12, space=8)
        centre("Artificial Intelligence and Data Engineering", 15, True, TODOCOL, 30)
        centre("Presented and defended by:", 12, space=8)
        centre(AUTHOR, 16, True, space=8)
        centre("On [Weekday, DD Month 2026]", 12, False, TODOCOL, 26)
        centre("Title", 12, space=8)
        centre(TITLE, 17, True, TODOCOL, 32)
        _label_block(d, [("Supervisor", ["Dr. Daoud Baalbaki"]),
                         ("Reviewers", ["Dr. [Reviewer 1]", "Dr. [Reviewer 2]"])])
        d.add_paragraph().paragraph_format.space_after = Pt(22)
        centre("Lebanese University – Faculty of Sciences", 12)


# ======================================================================
#  CONTENT
# ======================================================================
def build(variant="sciences", out=OUT):
    D = Doc()
    d = D.d

    title_page(D, variant)

    # ---------------- FRONT MATTER (roman numerals) ----------------
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    set_page_numbering(sec, "lowerRoman", start=1)
    add_footer_pagenum(sec)

    D.h1("Acknowledgements")
    D.todo("Personalise this. The reference thesis thanks, in order: the company supervisor, "
           "the university supervisor, the reviewers, project partners, then family. Keep to one page.")
    D.p("I would like to express my sincere gratitude to **Dr. Daoud Baalbaki** for his supervision "
        "throughout this project, and for the opportunity to carry out this work within **WhiteStork**, "
        "where I was able to apply what I learned in practice.")
    D.p("A special thanks goes to my university supervisor, **Dr. [Name]**, for his continuous guidance "
        "throughout my studies.")
    D.p("I would also like to thank the reviewers for their time and for the feedback that improved the "
        "quality of this work.")
    D.p("Finally, my appreciation extends to my family and friends for their constant encouragement and "
        "patience.")
    D.p("")
    D.p("*Rokaya Al Harakeh*")
    D.p("Beirut, [date]")

    D.h1("Abstract")
    D.p("Word Sense Disambiguation (WSD) — selecting the intended meaning of an ambiguous word from a "
        "predefined sense inventory — is a long-standing problem in Arabic natural language processing, "
        "made harder by templatic morphology, clitic attachment, and the routine omission of diacritics. "
        "Recent work has evaluated large generative language models on this task, but such models are "
        "expensive to adapt, and the reported evaluations rely on metrics whose behaviour on gloss-based "
        "Arabic datasets has not been examined.")
    D.p("This project investigates **parameter-efficient adaptation** of a small language model, "
        "**Gemma 2-2B**, under two distinct training objectives. First, supervised instruction fine-tuning "
        "(SFT) is applied to the El-Razzaz gloss-based Arabic WSD dataset using QLoRA, a 4-bit quantized "
        "low-rank adaptation scheme, on a single free-tier GPU. Second, continued pretraining (CPT) is "
        "applied to a Lebanese legal corpus under a causal language-modelling objective, and benchmarked "
        "against the untrained base model.")
    D.p("The fine-tuned 2B model attains **90.42% accuracy** on the 3,110-instance test set, matching "
        "published 8B-parameter models and exceeding the same-family 9B model, at a fraction of the "
        "parameter count and compute budget. Beyond replication, this work contributes an evaluation "
        "critique of the benchmark itself: it establishes that the dataset's trivial positional baseline "
        "is **64.95%** rather than the assumed 50%, that a classical lexical-overlap method adds "
        "essentially nothing over that baseline, and that the macro-averaged F1 reported in the literature "
        "is degenerate on this dataset — it reduces to accuracy rescaled by a denominator the model "
        "inflates through the variety of its own errors. Error analysis shows the fine-tuned model produced "
        "**zero malformed outputs** across the test set and exhibits no positional bias.")
    D.p("The continued-pretraining experiment adapts the same base model to a 27.2-million-token corpus "
        "of Lebanese legal Arabic in a single epoch. Held-out perplexity falls from **806.47 to 4.14** "
        "and next-token accuracy rises from 21.3% to **69.8%**, improving on every one of the seven "
        "corpus sources; a 300-item cloze probe constructed from held-out documents rises from 43.3% to "
        "**84.3%** under constrained scoring, including on statutory-instrument items where the base "
        "model performs no better than the majority baseline. Evaluated on the word-sense task it was "
        "never trained for, the adapted model **does not degrade** — accuracy rises from 24.9% to 36.5% "
        "— but decomposing that rise shows it to be **entirely a gain in format compliance** (52.5% to "
        "75.1%) with no change in discrimination (47.4% to 48.6%, not significant).")
    D.p("Taken together, the two experiments support a dissociation: **supervised fine-tuning changes task "
        "format, while continued pretraining changes domain distribution**. The two objectives are not "
        "interchangeable, and the distinction holds independently of effect size.")
    D.p("**Keywords:** Arabic NLP, Word Sense Disambiguation, Large Language Models, Parameter-Efficient "
        "Fine-Tuning, LoRA, QLoRA, Continued Pretraining, Gemma 2.")

    D.h1("List of Abbreviations and Symbols")
    D.table([
        ["Abbreviation", "Meaning"],
        ["BPE", "Byte Pair Encoding"],
        ["CLM", "Causal Language Modelling"],
        ["CPT", "Continued Pretraining"],
        ["DAPT", "Domain-Adaptive Pretraining"],
        ["GQA", "Grouped-Query Attention"],
        ["LLM", "Large Language Model"],
        ["LoRA", "Low-Rank Adaptation"],
        ["MSA", "Modern Standard Arabic"],
        ["NF4", "4-bit NormalFloat"],
        ["PEFT", "Parameter-Efficient Fine-Tuning"],
        ["PPL", "Perplexity"],
        ["QLoRA", "Quantized Low-Rank Adaptation"],
        ["SFT", "Supervised Fine-Tuning"],
        ["TAPT", "Task-Adaptive Pretraining"],
        ["WSD", "Word Sense Disambiguation"],
        ["r", "LoRA rank"],
        ["α", "LoRA scaling parameter"],
        ["W₀", "Frozen pretrained weight matrix"],
        ["ΔW", "Learned weight update"],
    ], widths=[4, 11])

    D.h1("Table of Contents")
    tp = d.add_paragraph()
    add_field(tp, 'TOC \\o "1-3" \\h \\z \\u',
              "Right-click here → Update Field → Update entire table.")

    D.h1("List of Tables")
    D.list_of("Table")

    D.h1("List of Figures")
    D.list_of("Figure")

    D.h1("List of Listings")
    D.list_of("Listing")
    D.todo("All four lists above are live Word fields. To build them: Ctrl+A, then F9, then choose "
           "'Update entire table' for each prompt. Repeat this once more immediately before printing, "
           "after all page numbers have settled.")

    # ---------------- BODY (arabic numerals) ----------------
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    set_page_numbering(sec, "decimal", start=1)
    add_footer_pagenum(sec)

    build_general_intro(D)
    build_ch1(D)
    build_ch2(D)
    build_ch3(D)
    build_ch4(D)
    build_ch5(D)
    build_ch6(D)
    build_ch7(D)
    build_ch8(D)
    build_ch9(D)
    build_ch10(D)
    build_ch11(D)
    build_conclusion(D)
    build_references(D)
    build_appendix_a(D)
    build_appendix_b(D)
    build_appendix_c(D)
    build_appendix_d(D)

    saved = D.save(out)
    print(f"Wrote {saved}")


# ---------------------------------------------------------------- intro
def build_general_intro(D):
    D.h1("General Introduction")
    D.p("Arabic is among the most widely spoken languages in the world, yet it remains under-served by "
        "modern language technology. One reason is structural: Arabic writing omits the short vowels that "
        "distinguish many words, so a single written form frequently corresponds to several unrelated "
        "meanings. Resolving which meaning is intended — *word sense disambiguation* — is therefore not a "
        "peripheral problem in Arabic but a prerequisite for reliable search, translation, and question "
        "answering.")
    D.p("Large generative language models have recently been applied to this task, with encouraging "
        "reported results. However, adapting such models is expensive: fine-tuning a model of even a few "
        "billion parameters by conventional means exceeds the memory of the hardware available to most "
        "students and small organisations. This raises a practical question that this project addresses "
        "directly: *how small a model, and how little compute, is actually required?*")
    D.p("The project pursues three objectives. The first is **replication under constraint**: to reproduce "
        "a published Arabic word sense disambiguation result using a model of roughly a quarter the size, "
        "on free, publicly available hardware, and to report honestly whether the smaller model holds up. "
        "The second is **evaluation criticism**: to establish what score the benchmark yields without a "
        "model at all, and to check whether its reported metrics measure what they appear to measure — "
        "questions the existing literature on this dataset does not ask. The third is **comparison of "
        "adaptation objectives**: to apply supervised fine-tuning and continued pretraining to the same "
        "base model and to characterise what each of them changes.")
    D.p("What was built is a four-stage pipeline that stages the dataset into an instruction-tuning "
        "format, adapts the base model with QLoRA, runs deterministic inference over the held-out test "
        "split, and scores the result — together with an offline analysis layer that reconstructs "
        "per-instance predictions and decomposes the errors. The entire pipeline is parameterised through "
        "environment variables and runs unmodified for every experiment reported here; its complete source "
        "is reproduced in Appendix A.")
    D.p("The contributions of this work are as follows. First, it demonstrates that a 2-billion-parameter "
        "model adapted with QLoRA on a single free-tier GPU matches published results from models four "
        "times larger. Second, it provides an evaluation critique of the benchmark used, establishing its "
        "true trivial baseline and demonstrating that its headline secondary metric is degenerate. Third, "
        "it contrasts two distinct adaptation objectives — supervised fine-tuning and continued "
        "pretraining — on the same model, and articulates what each one changes.")


def build_ch1(D):
    D.h1("Chapter 1 — Introduction")

    # ---- 1.1 -------------------------------------------------------
    D.h2("1.1  Motivation")
    D.p("Arabic is among the most widely spoken languages in the world, with several hundred million "
        "speakers, and yet in natural language processing it is routinely treated as low-resource. The "
        "shortfall is not one of raw text — Arabic is abundantly written — but of tooling, annotated "
        "resources, and models whose design accounts for how the language actually works.")
    D.p("A large part of the difficulty is structural. Arabic writing omits the short vowels that "
        "distinguish many otherwise identical words, so a single written form frequently corresponds to "
        "several unrelated meanings. The consonantal skeleton \\ar{كتب} may be read as *kataba* (he wrote), "
        "*kutiba* (it was written), or *kutub* (books). A human reader resolves this continuously and "
        "unconsciously by reading context; any system that processes Arabic must do the same. Resolving "
        "which of several possible meanings applies — **word sense disambiguation** — is therefore not a "
        "specialised subproblem in Arabic but a precondition for machine translation, search, and "
        "information retrieval to work at all.")
    D.p("Large generative language models have recently been applied to this task with encouraging "
        "results. But adapting them is expensive. Fine-tuning a model of even a few billion parameters by "
        "conventional means requires memory an order of magnitude beyond what is available to a student or "
        "a small organisation, as Chapter 4 quantifies. Parameter-efficient fine-tuning is the technique "
        "that closes this gap, making it possible to adapt a capable model on hardware that costs nothing.")
    D.p("That combination — a task that Arabic makes unavoidable, and a method that makes adaptation "
        "affordable — motivates the question this report addresses: **how small a model, and how little "
        "compute, does this task actually require?**")

    # ---- 1.2 -------------------------------------------------------
    expand_1_1(D)

    D.h2("1.2  Problem statement")
    D.p("The task is defined as follows. Given a sentence s, a target word w occurring in s, and a "
        "candidate sense set C = {c₁, …, c_k} where each candidate is supplied with an Arabic gloss, the "
        "system must predict the identifier of the sense that matches the use of w in s.")
    D.p("In the benchmark used here the candidate set always contains exactly two entries, so the task is "
        "a binary choice between two visible definitions. This is a consequence of how the dataset was "
        "constructed, and it turns out to be the single most consequential structural fact in the entire "
        "study; it is examined in Section 3.6 and again in the results.")
    D.p("Two research questions follow.")
    D.numbered("**RQ1.** Can a model below three billion parameters match the performance of published "
               "seven- to nine-billion-parameter models on Arabic word sense disambiguation, when trained "
               "under a free-tier compute budget?")
    D.numbered("**RQ2.** What does each adaptation objective actually change? Specifically, how does "
               "supervised instruction fine-tuning differ from continued pretraining in its effect on the "
               "same model?")
    D.p("The first question is empirical and is answered by replication at reduced scale. The second is "
        "conceptual, and is answered by running both objectives on one model and measuring what each one "
        "moves.")

    # ---- 1.3 -------------------------------------------------------
    D.h2("1.3  Objectives and scope")
    D.p("The project has four objectives:")
    D.bullet("Reproduce the Gemma result of the replicated study using Gemma 2-2B in place of Gemma 2-9B, "
             "under QLoRA on a single free-tier GPU.")
    D.bullet("Establish the trivial and non-neural baselines that the published literature on this dataset "
             "does not report, so that the reported accuracies can be interpreted.")
    D.bullet("Audit the evaluation metrics used by the benchmark, and determine whether they carry the "
             "information they are assumed to carry.")
    D.bullet("Apply continued pretraining to a Lebanese legal corpus under a causal language-modelling "
             "objective, and contrast its effect with that of supervised fine-tuning.")
    D.p("The scope is deliberately narrow, and the boundaries are stated here so that the results are read "
        "correctly.")
    D.table([
        ["Dimension", "Scope of this work", "Reason"],
        ["Dataset", "Dataset A (El-Razzaz) only", "Dataset B (SALMA) reaches ~4,096-token sequences; "
         "combined with Arabic tokenizer fertility this exceeds the available memory budget "
         "(Sections 2.3, 3.6)"],
        ["Adaptation method", "QLoRA only; base weights never modified", "Full fine-tuning requires "
         "~41.7 GB against 16 GB available (Section 4.1)"],
        ["Model", "Gemma 2-2B, single family and size", "Matches the replicated study's family; scaled "
         "down to test the capacity question"],
        ["Runs", "Single seed, no variance estimate", "Compute budget; recorded as a limitation"],
        ["Language variety", "Modern Standard Arabic", "Inherited from the dataset; dialects out of scope"],
    ], widths=[3.2, 4.4, 6.4])
    D.caption("Scope of the study, and the reason for each boundary.")

    # ---- 1.4 -------------------------------------------------------
    D.h2("1.4  Contributions")
    D.p("This work makes six contributions.")
    D.numbered("**Performance at reduced scale.** A 2-billion-parameter model adapted with QLoRA on a "
               "single free-tier GPU reaches 90.42% accuracy on the 3,110-instance test set — matching the "
               "published 8-billion-parameter result exactly and exceeding the same-family 9-billion-"
               "parameter model, at roughly a quarter of the parameters.")
    D.numbered("**A trivial baseline the literature omits.** Because every instance offers exactly two "
               "candidates and the correct answer is the second one about 65% of the time, always choosing "
               "the second candidate achieves **64.95%** without any model. This baseline is stable across "
               "all three data splits, establishing it as a property of dataset construction rather than "
               "an artefact of sampling. Every published accuracy on this benchmark should be read against "
               "it rather than against 50%.")
    D.numbered("**A non-neural baseline.** Simplified Lesk, the classical lexical-overlap method, reaches "
               "64.98% — an improvement of 0.03 points over the trivial rule. Lexical overlap is therefore "
               "not a viable strategy on this data, which is what makes the neural result meaningful.")
    D.numbered("**A demonstration that the secondary metric is degenerate.** Every gold class in the test "
               "set has support of exactly one. Macro-F1 consequently reduces to accuracy rescaled by a "
               "denominator that the model inflates through the variety of its own errors, and carries no "
               "information beyond accuracy. The arithmetic is exact and the same structure is visible in "
               "the published figures.")
    D.numbered("**An error analysis.** The fine-tuned model produced zero malformed or out-of-set outputs "
               "across all 3,110 items, so every error is genuine sense confusion rather than a formatting "
               "failure. It also shows no positional bias: it selects the second candidate less often than "
               "the data does, and is more accurate on the rarer first position.")
    D.numbered("**A measurement of tokenizer behaviour on Arabic.** The Gemma 2 tokenizer requires 1.79 "
               "times more subword tokens per word on Arabic than on English, measured on this corpus. "
               "This bounds the effective context, explains the exclusion of the larger benchmark, and "
               "anticipates the behaviour of specialised legal Arabic.")
    D.p("Taken together, the last four contributions support the central interpretive claim of the report: "
        "on this benchmark the model is asked to discriminate between two definitions that are already "
        "present in its input, not to retrieve knowledge from its parameters. That is why capacity buys so "
        "little, and it is the basis for the dissociation drawn in Chapter 10.")

    # ---- 1.5 -------------------------------------------------------
    D.h2("1.5  Report structure")
    D.p("The remainder of this report is organised as follows.")
    D.bullet("**Chapter 2** establishes the architectural background: the attention mechanism that makes "
             "contextual disambiguation possible, the specific configuration of Gemma 2-2B, a measurement "
             "of tokenizer fertility on Arabic, and the causal language-modelling objective.")
    D.bullet("**Chapter 3** covers the linguistic sources of Arabic ambiguity, the orthographic decisions "
             "taken here, the word sense disambiguation task and its resources, and the evaluation metrics.")
    D.bullet("**Chapter 4** develops the adaptation methodology from full fine-tuning through the "
             "parameter-efficient landscape to LoRA, quantization and QLoRA, closing with the distinction "
             "between training objectives on which the conclusion depends.")
    D.bullet("**Chapter 5** reviews related work, describes the replicated study in detail, and states the "
             "gap this report addresses.")
    D.bullet("**Chapter 6** documents the supervised fine-tuning methodology: data staging, prompt "
             "construction, configuration, training, inference and evaluation, together with the "
             "deviations from the replicated recipe and the implementation problems encountered.")
    D.bullet("**Chapter 7** reports the results: the headline accuracy and its comparison against "
             "published models, the trivial and non-neural baselines established for this dataset, "
             "the demonstration that the standard secondary metric is degenerate on it, a "
             "decomposition of the remaining errors, and the interpretation that follows.")
    D.bullet("**Chapter 8** documents the continued pre-training methodology: the legal corpus and its "
             "composition, the document-level split and the contamination it prevents, packing, the "
             "training configuration set against the supervised one row by row, the evaluation design, "
             "and the predictions registered before the run.")
    D.bullet("**Chapter 9** reports the continued pre-training results: the training curve, perplexity "
             "and next-token accuracy broken down by corpus source, the legal cloze probe under two "
             "scoring regimes, the retention test on Dataset A and its decomposition, and the "
             "pre-registered predictions scored against what happened.")
    D.bullet("**Chapter 10** draws the two halves together into the dissociation that is the conclusion "
             "of this project: supervised fine-tuning changes task format, continued pre-training "
             "changes domain distribution, and the two are not substitutes.")
    D.bullet("**Chapter 11** states the limitations of both experiments in full.")
    D.bullet("**Appendix A** contains the complete pipeline source code and **Appendix B** the "
             "supervised run configuration recovered from the training artifacts; **Appendix C** does "
             "the same for the continued pre-training run, and **Appendix D** reproduces the Arabic "
             "generation samples.")


if __name__ == "__main__":
    # chapter builders live in build_chapters.py to keep this file readable
    from build_chapters import (build_ch2, build_ch3, build_ch4, build_ch5,
                                build_ch6, build_ch7, build_references,
                                expand_1_1)
    from build_appendices import build_conclusion, build_appendix_a, build_appendix_b
    # Chapters 8-11 -- the continued pre-training half
    from build_cpt_chapters import build_ch8, build_ch9, build_ch10, build_ch11
    from build_cpt_appendices import build_appendix_c, build_appendix_d
    globals().update(dict(build_ch2=build_ch2, build_ch3=build_ch3, build_ch4=build_ch4,
                          build_ch5=build_ch5, build_ch6=build_ch6, build_ch7=build_ch7,
                          build_ch8=build_ch8, build_ch9=build_ch9,
                          build_ch10=build_ch10, build_ch11=build_ch11,
                          expand_1_1=expand_1_1,
                          build_references=build_references,
                          build_conclusion=build_conclusion,
                          build_appendix_a=build_appendix_a,
                          build_appendix_b=build_appendix_b,
                          build_appendix_c=build_appendix_c,
                          build_appendix_d=build_appendix_d))
    # Two copies of the identical report, differing only in the title page.
    build("sciences", OUT)
    build("engineering", OUT2)
